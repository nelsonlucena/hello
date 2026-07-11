"""
Editorial Memo deck-kit: a small python-pptx component library implementing
the design system in references/style-guide.md.

This module is deliberately a *component library*, not a fixed set of
slide templates. build_deck.py composes these primitives into the five
reference slide types, but any new slide should still be built out of these
same functions so every deck produced by the skill stays visually
consistent. Do not hand-roll shapes/colors that bypass these helpers.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ---------------------------------------------------------------------------
# Slide geometry
# ---------------------------------------------------------------------------
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.55)

TOP_BAND_H = Inches(0.35)
BOTTOM_BAND_H = Inches(0.55)

# ---------------------------------------------------------------------------
# Fonts
# ---------------------------------------------------------------------------
SERIF = "Libre Baskerville"
SANS = "Source Sans 3"

# ---------------------------------------------------------------------------
# Palette (see references/style-guide.md — do not add colors here)
# ---------------------------------------------------------------------------
PAPER = RGBColor(0xFA, 0xF8, 0xF4)
INK = RGBColor(0x1A, 0x1A, 0x17)
BODY_GRAY = RGBColor(0x3C, 0x3C, 0x37)
MUTED_GRAY = RGBColor(0x79, 0x79, 0x72)
HAIRLINE = RGBColor(0xDD, 0xDC, 0xD8)
CARD_BORDER = RGBColor(0xC8, 0xC4, 0xB8)
BAND_BLACK = RGBColor(0x12, 0x12, 0x12)
OXBLOOD = RGBColor(0x8B, 0x1A, 0x1A)
STATUS_RED = RGBColor(0xC6, 0x28, 0x28)
STATUS_GREEN = RGBColor(0x2E, 0x7D, 0x32)
STATUS_AMBER = RGBColor(0xB0, 0x79, 0x00)

DARK_BG = RGBColor(0x16, 0x14, 0x0F)
CREAM_TEXT = RGBColor(0xF2, 0xEF, 0xE6)
TAN_TEXT = RGBColor(0xA2, 0x9D, 0x8E)
TERRACOTTA = RGBColor(0xD9, 0x8A, 0x59)
DARK_HAIRLINE = RGBColor(0x3A, 0x35, 0x2B)  # hairline rule on the dark closing slide

STATUS_COLORS = {
    "ON TRACK": STATUS_GREEN,
    "AT RISK": STATUS_RED,
    "PROGRAM-LEVEL": STATUS_RED,
    "EXECUTION": STATUS_AMBER,
}

# ---------------------------------------------------------------------------
# Type scale (see references/style-guide.md)
# ---------------------------------------------------------------------------
SZ_MICRO_FOOTER = Pt(8)
SZ_EYEBROW = Pt(10.5)
SZ_CTA_LINK = Pt(10.5)
SZ_CARD_LABEL = Pt(10.5)
SZ_BODY = Pt(12)
SZ_STAT_UNIT = Pt(12)
SZ_SUBTITLE = Pt(15)
SZ_DECISION_ITEM_TITLE = Pt(16)
SZ_H3 = Pt(18)
SZ_H2 = Pt(28)
SZ_H1 = Pt(36)
SZ_HERO = Pt(34)
SZ_HERO_XL = Pt(44)
SZ_DECISION_NUMBER = Pt(40)


# ---------------------------------------------------------------------------
# Low-level helpers
# ---------------------------------------------------------------------------

def div(emu_value, n):
    """Integer-safe division of an EMU length. python-pptx/OOXML require
    integer EMU attributes — plain `/` produces a float that corrupts the
    file (LibreOffice/PowerPoint will then refuse to open it)."""
    return Emu(int(emu_value) // n)


def new_presentation():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def add_blank_slide(prs):
    blank = prs.slide_layouts[6]
    return prs.slides.add_slide(blank)


def set_background(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_rect(slide, left, top, width, height, fill_color=None, line_color=None,
             line_weight=Pt(0.75)):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.shadow.inherit = False
    if fill_color is not None:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
    else:
        shape.fill.background()
    if line_color is not None:
        shape.line.color.rgb = line_color
        shape.line.width = line_weight
    else:
        shape.line.fill.background()
    return shape


def add_line(slide, left, top, width, height, color=HAIRLINE, weight=Pt(0.75)):
    """A hairline rule. Use width>0,height=0 for horizontal; width=0,height>0 for vertical."""
    conn = slide.shapes.add_connector(1, left, top, left + width, top + height)
    conn.shadow.inherit = False
    conn.line.color.rgb = color
    conn.line.width = weight
    return conn


def add_dot(slide, center_x, center_y, diameter, color):
    left = center_x - diameter // 2
    top = center_y - diameter // 2
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, diameter, diameter)
    dot.shadow.inherit = False
    dot.fill.solid()
    dot.fill.fore_color.rgb = color
    dot.line.fill.background()
    return dot


def _set_tracking(run, points):
    """Set letter-spacing on a run (python-pptx has no high-level API for this)."""
    rPr = run.font._rPr if run.font._rPr is not None else run._r.get_or_add_rPr()
    rPr.set('spc', str(int(points * 100)))


def add_textbox(slide, left, top, width, height, anchor=MSO_ANCHOR.TOP,
                 align=PP_ALIGN.LEFT, word_wrap=True):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = word_wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.paragraphs[0].alignment = align
    return box


def add_run(paragraph, text, font=SANS, size=SZ_BODY, color=INK, bold=False,
            italic=False, tracking=None):
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = size
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    if tracking is not None:
        _set_tracking(run, tracking)
    return run


def add_paragraph(text_frame, first=False):
    if first and len(text_frame.paragraphs) == 1 and not text_frame.paragraphs[0].runs:
        return text_frame.paragraphs[0]
    return text_frame.add_paragraph()


def set_line_spacing(paragraph, spacing_pt=None, space_after=None):
    if spacing_pt is not None:
        paragraph.line_spacing = spacing_pt
    if space_after is not None:
        paragraph.space_after = space_after


# ---------------------------------------------------------------------------
# Component: frame bands (paper slides)
# ---------------------------------------------------------------------------

def add_frame_bands(slide):
    add_rect(slide, 0, 0, SLIDE_W, TOP_BAND_H, fill_color=BAND_BLACK)
    add_rect(slide, 0, SLIDE_H - BOTTOM_BAND_H, SLIDE_W, BOTTOM_BAND_H, fill_color=BAND_BLACK)


# ---------------------------------------------------------------------------
# Component: eyebrow (plain or breadcrumb, with hairline continuation)
# ---------------------------------------------------------------------------

def add_eyebrow(slide, left, top, width, segments, with_rule=True, rule_to=None):
    """segments: list of (text, color) tuples rendered in sequence, all caps + tracked."""
    box = add_textbox(slide, left, top, width, Inches(0.25))
    p = box.text_frame.paragraphs[0]
    for i, (text, color) in enumerate(segments):
        prefix = "   " if i > 0 else ""
        add_run(p, prefix + text.strip().upper(), font=SANS, size=SZ_EYEBROW, color=color,
                bold=True, tracking=1.5)
    if with_rule:
        rule_left = left + Inches(2.7)
        rule_right = rule_to if rule_to is not None else (SLIDE_W - MARGIN)
        add_line(slide, rule_left, top + Inches(0.12), rule_right - rule_left, 0,
                 color=HAIRLINE, weight=Pt(0.75))
    return box


# ---------------------------------------------------------------------------
# Component: serif title with exactly one italic phrase
# ---------------------------------------------------------------------------

def add_title(slide, left, top, width, height, lines, size=SZ_H2,
              ink_color=INK, accent_color=OXBLOOD, align=PP_ALIGN.LEFT,
              line_spacing=1.05):
    """
    lines: list of paragraphs, each a list of (text, is_italic) tuples.
    Exactly one (text, True) run across the whole title should be the accent phrase.
    """
    box = add_textbox(slide, left, top, width, height, align=align)
    tf = box.text_frame
    for i, line in enumerate(lines):
        p = add_paragraph(tf, first=(i == 0))
        p.alignment = align
        set_line_spacing(p, spacing_pt=line_spacing)
        for text, italic in line:
            color = accent_color if italic else ink_color
            add_run(p, text, font=SERIF, size=size, color=color, bold=not italic,
                     italic=italic)
    return box


# ---------------------------------------------------------------------------
# Component: status dot pill
# ---------------------------------------------------------------------------

def add_status_pill(slide, right_x, center_y, label, size=SZ_CARD_LABEL):
    """Right-aligned status pill: dot + label, dot appears left of label, whole
    group ends at right_x."""
    color = STATUS_COLORS.get(label.upper(), MUTED_GRAY)
    label_text = label.upper()
    # estimate label width to place the dot; simplest robust approach is a
    # single textbox with two runs anchored by right edge via a wide box.
    box_w = Inches(1.6)
    box = add_textbox(slide, right_x - box_w, center_y - Inches(0.09), box_w, Inches(0.2),
                       align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    p = box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    add_run(p, "●  " + label_text, font=SANS, size=size, color=color,
            bold=True, tracking=1.0)
    return box


# ---------------------------------------------------------------------------
# Component: hero number + unit label
# ---------------------------------------------------------------------------

def add_hero_number(slide, left, top, number_text, unit_text, number_size=SZ_HERO,
                     unit_size=SZ_STAT_UNIT, color=INK, unit_color=MUTED_GRAY,
                     width=Inches(4), height=Inches(0.85)):
    box = add_textbox(slide, left, top, width, height)
    p = box.text_frame.paragraphs[0]
    add_run(p, number_text + " ", font=SERIF, size=number_size, color=color, bold=True)
    if unit_text:
        add_run(p, unit_text.upper(), font=SERIF, size=unit_size, color=unit_color,
                bold=False)
    return box


# ---------------------------------------------------------------------------
# Component: body text block (supports bold emphasis spans)
# ---------------------------------------------------------------------------

def add_body_text(slide, left, top, width, height, spans, size=SZ_BODY,
                   color=BODY_GRAY, align=PP_ALIGN.LEFT, line_spacing=1.15,
                   anchor=MSO_ANCHOR.TOP):
    """spans: list of paragraphs; each paragraph is a list of (text, bold) tuples."""
    box = add_textbox(slide, left, top, width, height, align=align, anchor=anchor)
    tf = box.text_frame
    for i, para in enumerate(spans):
        p = add_paragraph(tf, first=(i == 0))
        p.alignment = align
        set_line_spacing(p, spacing_pt=line_spacing)
        for text, bold in para:
            add_run(p, text, font=SANS, size=size, color=color, bold=bold)
    return box


def add_bullets(slide, left, top, width, height, items, size=SZ_BODY,
                 color=BODY_GRAY, bullet_color=OXBLOOD, line_spacing=1.15,
                 space_after=Pt(8)):
    """items: list of paragraphs, each a list of (text, bold) run tuples."""
    box = add_textbox(slide, left, top, width, height)
    tf = box.text_frame
    for i, para in enumerate(items):
        p = add_paragraph(tf, first=(i == 0))
        set_line_spacing(p, spacing_pt=line_spacing, space_after=space_after)
        add_run(p, "●  ", font=SANS, size=size, color=bullet_color, bold=False)
        for text, bold in para:
            add_run(p, text, font=SANS, size=size, color=color, bold=bold)
    return box


# ---------------------------------------------------------------------------
# Component: footer
# ---------------------------------------------------------------------------

def add_footer(slide, deck_name, page_number, dark=False):
    color = TAN_TEXT if dark else MUTED_GRAY
    top = SLIDE_H - BOTTOM_BAND_H + Inches(0.18) if dark else SLIDE_H - BOTTOM_BAND_H - Inches(0.32)
    left_box = add_textbox(slide, MARGIN, top, Inches(6), Inches(0.25))
    p = left_box.text_frame.paragraphs[0]
    add_run(p, deck_name.upper(), font=SANS, size=SZ_MICRO_FOOTER, color=color,
            bold=False, tracking=1.5)

    right_box = add_textbox(slide, SLIDE_W - MARGIN - Inches(1.5), top, Inches(1.5),
                             Inches(0.25), align=PP_ALIGN.RIGHT)
    p2 = right_box.text_frame.paragraphs[0]
    p2.alignment = PP_ALIGN.RIGHT
    add_run(p2, str(page_number).zfill(2), font=SANS, size=SZ_MICRO_FOOTER, color=color,
            bold=False, tracking=1.5)


# ---------------------------------------------------------------------------
# Component: card container
# ---------------------------------------------------------------------------

def add_card(slide, left, top, width, height, border_color=HAIRLINE,
             left_accent_color=None, fill_color=None):
    add_rect(slide, left, top, width, height, fill_color=fill_color, line_color=border_color,
              line_weight=Pt(0.75))
    if left_accent_color is not None:
        add_rect(slide, left, top, Pt(3), height, fill_color=left_accent_color)
    return None


# ---------------------------------------------------------------------------
# Component: CTA link line
# ---------------------------------------------------------------------------

def add_cta(slide, left, top, text, color=OXBLOOD):
    box = add_textbox(slide, left, top, Inches(4), Inches(0.25))
    p = box.text_frame.paragraphs[0]
    add_run(p, text.upper() + "  →", font=SANS, size=SZ_CTA_LINK, color=color,
            bold=True, tracking=1.0)
    return box
