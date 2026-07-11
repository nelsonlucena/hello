#!/usr/bin/env python3
"""
Build an "Editorial Memo" style PowerPoint deck from a JSON content spec.

Usage:
    python3 build_deck.py content.json output.pptx

See references/style-guide.md for the design system this implements, and
assets/example_content.json for the content-spec schema (by example).
Only compose slides out of deck_kit.py's components — see that module's
docstring before adding a new slide type.
"""
import json
import sys
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

import deck_kit as dk

MARGIN = dk.MARGIN
CONTENT_TOP = dk.TOP_BAND_H + Inches(0.28)
CONTENT_RIGHT = dk.SLIDE_W - MARGIN
CONTENT_WIDTH = dk.SLIDE_W - 2 * MARGIN


# ---------------------------------------------------------------------------
# JSON -> deck_kit segment helpers
# ---------------------------------------------------------------------------

def _seg(item, key):
    """Normalize a segment: plain string, or {"text":..., key: true}."""
    if isinstance(item, str):
        return (item, False)
    return (item["text"], bool(item.get(key, False)))


def title_lines(spec):
    """spec: list of lines, each a list of strings / {"text","italic"}."""
    return [[_seg(seg, "italic") for seg in line] for line in spec]


def body_paras(spec):
    """spec: list of paragraphs, each a list of strings / {"text","bold"}."""
    return [[_seg(seg, "bold") for seg in para] for para in spec]


def eyebrow_segments(spec, default_color=dk.MUTED_GRAY, accent_color=dk.OXBLOOD):
    """spec: list of strings / {"text","accent":true}."""
    out = []
    for seg in spec:
        if isinstance(seg, str):
            out.append((seg, default_color))
        else:
            out.append((seg["text"], accent_color if seg.get("accent") else default_color))
    return out


# ---------------------------------------------------------------------------
# Slide builders
# ---------------------------------------------------------------------------

def build_cover(prs, data, deck_name):
    slide = dk.add_blank_slide(prs)
    dk.set_background(slide, dk.PAPER)
    dk.add_frame_bands(slide)

    top = dk.TOP_BAND_H + Inches(0.22)

    # memo header
    memo = data.get("memo")
    conf_label = data.get("confidential_label")
    if conf_label:
        box = dk.add_textbox(slide, MARGIN, top, Inches(6), Inches(0.22))
        p = box.text_frame.paragraphs[0]
        dk.add_run(p, conf_label.upper(), font=dk.SANS, size=dk.SZ_EYEBROW,
                   color=dk.OXBLOOD, bold=True, tracking=1.5)
    if memo:
        rows = [("To", memo.get("to", "")), ("From", memo.get("from", "")),
                ("Date", memo.get("date", ""))]
        row_h = Inches(0.2)
        for i, (label, value) in enumerate(rows):
            box = dk.add_textbox(slide, dk.SLIDE_W - MARGIN - Inches(4),
                                  top + row_h * i, Inches(4), row_h,
                                  align=PP_ALIGN.RIGHT)
            p = box.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.RIGHT
            dk.add_run(p, label + "  ", font=dk.SANS, size=dk.SZ_BODY, color=dk.INK,
                       bold=True)
            dk.add_run(p, value, font=dk.SANS, size=dk.SZ_BODY, color=dk.BODY_GRAY)

    y = top + Inches(0.62)
    eyebrow = data.get("eyebrow")
    if eyebrow:
        dk.add_eyebrow(slide, MARGIN, y, CONTENT_WIDTH, [(eyebrow, dk.MUTED_GRAY)],
                        with_rule=False)
    y += Inches(0.3)

    dk.add_title(slide, MARGIN, y, CONTENT_WIDTH, Inches(1.5),
                 title_lines(data["title"]), size=dk.SZ_H1)
    y += Inches(1.3) if len(data["title"]) > 1 else Inches(0.75)

    subtitle = data.get("subtitle")
    if subtitle:
        dk.add_body_text(slide, MARGIN, y, CONTENT_WIDTH, Inches(0.7),
                          [[(subtitle, False)]], size=dk.SZ_SUBTITLE,
                          color=dk.BODY_GRAY, line_spacing=1.2)
        y += Inches(0.75)

    cards = data.get("stat_cards", [])
    if cards:
        y += Inches(0.2)
        card_h = Inches(1.55)
        card_w = dk.div(CONTENT_WIDTH, len(cards))
        for i, card in enumerate(cards):
            cx = MARGIN + card_w * i
            dk.add_card(slide, cx, y, card_w, card_h, border_color=dk.HAIRLINE)
            pad = Inches(0.18)
            label_top = y + Inches(0.14)
            box = dk.add_textbox(slide, cx + pad, label_top, card_w - 2 * pad, Inches(0.2))
            p = box.text_frame.paragraphs[0]
            dk.add_run(p, card["label"].upper(), font=dk.SANS, size=dk.SZ_CARD_LABEL,
                       color=dk.INK, bold=True, tracking=1.0)
            if card.get("status"):
                dk.add_status_pill(slide, cx + card_w - pad, label_top + Inches(0.1),
                                    card["status"])
            if card.get("number"):
                dk.add_hero_number(slide, cx + pad, label_top + Inches(0.28),
                                    card["number"], card.get("unit", ""),
                                    number_size=dk.SZ_HERO, unit_size=dk.SZ_STAT_UNIT,
                                    height=Inches(0.55))
                body_top = label_top + Inches(0.94)
            else:
                body_top = label_top + Inches(0.32)
            if card.get("body"):
                dk.add_body_text(slide, cx + pad, body_top,
                                  card_w - 2 * pad, Inches(0.9),
                                  body_paras(card["body"]), size=Pt(10.5),
                                  line_spacing=1.15)
    dk.add_footer(slide, deck_name, data.get("page", 1))
    return slide


def build_progress(prs, data, deck_name, page):
    slide = dk.add_blank_slide(prs)
    dk.set_background(slide, dk.PAPER)
    dk.add_frame_bands(slide)

    y = dk.TOP_BAND_H + Inches(0.22)
    dk.add_eyebrow(slide, MARGIN, y, CONTENT_WIDTH,
                   eyebrow_segments(data["breadcrumb"]))
    y += Inches(0.32)
    dk.add_title(slide, MARGIN, y, CONTENT_WIDTH, Inches(0.65),
                 title_lines(data["title"]), size=dk.SZ_H2)
    y += Inches(0.62)
    dk.add_line(slide, MARGIN, y, CONTENT_WIDTH, 0, color=dk.HAIRLINE, weight=Pt(0.75))
    y += Inches(0.18)

    stats = data.get("stats", [])
    if stats:
        stat_w = dk.div(CONTENT_WIDTH, len(stats))
        for i, s in enumerate(stats):
            sx = MARGIN + stat_w * i
            dk.add_hero_number(slide, sx, y, s["number"], s.get("unit", ""),
                                number_size=dk.SZ_HERO_XL, unit_size=dk.SZ_STAT_UNIT,
                                height=Inches(0.68))
            cap_y = y + Inches(0.78)
            spans = []
            if s.get("caption_bold"):
                spans.append((s["caption_bold"] + " ", True))
            if s.get("caption"):
                spans.append((s["caption"], False))
            dk.add_body_text(slide, sx, cap_y, stat_w - Inches(0.25), Inches(0.5),
                              [spans], size=Pt(10.5), line_spacing=1.15)
            if i > 0:
                dk.add_line(slide, sx - Inches(0.0), y, 0, Inches(0.9), color=dk.HAIRLINE)
        y += Inches(1.3)
        dk.add_line(slide, MARGIN, y, CONTENT_WIDTH, 0, color=dk.HAIRLINE, weight=Pt(0.75))
        y += Inches(0.22)

    columns = data.get("columns", [])
    if columns:
        col_w = dk.div(CONTENT_WIDTH - Inches(0.4), len(columns))
        for i, col in enumerate(columns):
            cx = MARGIN + (col_w + Inches(0.4)) * i
            if i > 0:
                dk.add_line(slide, cx - Inches(0.2), y, 0, Inches(2.6), color=dk.HAIRLINE)
            heading_box = dk.add_textbox(slide, cx, y, col_w, Inches(0.22))
            p = heading_box.text_frame.paragraphs[0]
            dk.add_run(p, col["heading"].upper(), font=dk.SANS, size=dk.SZ_EYEBROW,
                       color=dk.OXBLOOD, bold=True, tracking=1.0)
            cy = y + Inches(0.3)
            if col.get("bullets"):
                dk.add_bullets(slide, cx, cy, col_w, Inches(2.2),
                                body_paras(col["bullets"]), size=Pt(11))
            if col.get("inline_stats"):
                sw = dk.div(col_w, len(col["inline_stats"]))
                for j, s in enumerate(col["inline_stats"]):
                    num_color = dk.OXBLOOD if s.get("color") == "oxblood" else dk.INK
                    dk.add_hero_number(slide, cx + sw * j, cy, s["number"], "",
                                        number_size=dk.SZ_HERO,
                                        color=num_color, width=sw - Inches(0.1))
                    lbl_box = dk.add_textbox(slide, cx + sw * j, cy + Inches(0.55),
                                              sw - Inches(0.1), Inches(0.4))
                    lp = lbl_box.text_frame.paragraphs[0]
                    for k, line in enumerate(s.get("label", "").split("\n")):
                        lp2 = lp if k == 0 else lbl_box.text_frame.add_paragraph()
                        dk.add_run(lp2, line.upper(), font=dk.SANS, size=Pt(8.5),
                                   color=dk.MUTED_GRAY, bold=True, tracking=0.8)
                cy += Inches(1.0)
            if col.get("body"):
                dk.add_body_text(slide, cx, cy, col_w, Inches(1.2),
                                  body_paras(col["body"]), size=Pt(10.5),
                                  line_spacing=1.2)

    dk.add_footer(slide, deck_name, page)
    return slide


def build_two_column(prs, data, deck_name, page):
    slide = dk.add_blank_slide(prs)
    dk.set_background(slide, dk.PAPER)
    dk.add_frame_bands(slide)

    y = dk.TOP_BAND_H + Inches(0.22)
    dk.add_eyebrow(slide, MARGIN, y, CONTENT_WIDTH, eyebrow_segments(data["breadcrumb"]))
    y += Inches(0.32)
    dk.add_title(slide, MARGIN, y, CONTENT_WIDTH, Inches(0.65),
                 title_lines(data["title"]), size=dk.SZ_H2)
    y += Inches(0.72)

    columns = data["columns"]
    col_gap = Inches(0.5)
    col_w = dk.div(CONTENT_WIDTH - col_gap, len(columns))
    col_top = y
    for i, col in enumerate(columns):
        cx = MARGIN + (col_w + col_gap) * i
        if i > 0:
            dk.add_line(slide, cx - dk.div(col_gap, 2), col_top, 0, Inches(3.9), color=dk.HAIRLINE)
        name_box = dk.add_textbox(slide, cx, col_top, col_w - Inches(1.1), Inches(0.35))
        p = name_box.text_frame.paragraphs[0]
        dk.add_run(p, col["name"], font=dk.SERIF, size=Pt(18), color=dk.INK, bold=True)
        if col.get("status"):
            dk.add_status_pill(slide, cx + col_w, col_top + Inches(0.12), col["status"])
        cy = col_top + Inches(0.4)
        if col.get("subtitle"):
            dk.add_body_text(slide, cx, cy, col_w, Inches(0.3),
                              [[(col["subtitle"], False)]], size=Pt(10.5),
                              color=dk.MUTED_GRAY)
            cy += Inches(0.35)
        if col.get("number"):
            box = dk.add_textbox(slide, cx, cy, col_w, Inches(1.0))
            p2 = box.text_frame.paragraphs[0]
            p2.line_spacing = 1.05
            dk.add_run(p2, col["number"] + " ", font=dk.SERIF, size=dk.SZ_HERO_XL,
                       color=dk.INK, bold=True)
            if col.get("unit"):
                dk.add_run(p2, col["unit"].upper() + "  ", font=dk.SERIF,
                           size=dk.SZ_STAT_UNIT, color=dk.MUTED_GRAY)
            if col.get("lead"):
                lead_spans = body_paras([col["lead"]])[0] if isinstance(col["lead"], list) else [(col["lead"], False)]
                for text, bold in lead_spans:
                    dk.add_run(p2, text, font=dk.SANS, size=Pt(11), color=dk.BODY_GRAY,
                               bold=bold)
            cy += Inches(1.05)
        if col.get("bullets"):
            dk.add_bullets(slide, cx, cy, col_w, Inches(2.0), body_paras(col["bullets"]),
                            size=Pt(11))

    dk.add_footer(slide, deck_name, page)
    return slide


def build_risks(prs, data, deck_name, page):
    slide = dk.add_blank_slide(prs)
    dk.set_background(slide, dk.PAPER)
    dk.add_frame_bands(slide)

    y = dk.TOP_BAND_H + Inches(0.22)
    dk.add_eyebrow(slide, MARGIN, y, CONTENT_WIDTH, eyebrow_segments(data["breadcrumb"]))
    y += Inches(0.32)
    dk.add_title(slide, MARGIN, y, CONTENT_WIDTH, Inches(0.65),
                 title_lines(data["title"]), size=dk.SZ_H2)
    y += Inches(0.75)

    cards = data["cards"]
    gap = Inches(0.35)
    card_w = dk.div(CONTENT_WIDTH - gap * (len(cards) - 1), len(cards))
    card_h = Inches(4.15)
    for i, card in enumerate(cards):
        cx = MARGIN + (card_w + gap) * i
        accent = dk.STATUS_RED if card.get("status", "").upper() in ("AT RISK", "PROGRAM-LEVEL") else dk.OXBLOOD
        dk.add_card(slide, cx, y, card_w, card_h, border_color=dk.HAIRLINE,
                    left_accent_color=accent)
        pad = Inches(0.24)
        header_box = dk.add_textbox(slide, cx + pad, y + Inches(0.16), card_w - 2 * pad, Inches(0.2))
        p = header_box.text_frame.paragraphs[0]
        dk.add_run(p, card["label"].upper(), font=dk.SANS, size=dk.SZ_CARD_LABEL,
                   color=dk.INK, bold=True, tracking=1.0)
        if card.get("status"):
            dk.add_status_pill(slide, cx + card_w - pad, y + Inches(0.26), card["status"])
        dk.add_title(slide, cx + pad, y + Inches(0.42), card_w - 2 * pad, Inches(0.6),
                     title_lines(card["title"]), size=dk.SZ_H3)
        dk.add_body_text(slide, cx + pad, y + Inches(1.05), card_w - 2 * pad, Inches(2.3),
                          body_paras(card["body"]), size=Pt(11), line_spacing=1.25)
        gate_y = y + card_h - Inches(0.75)
        dk.add_line(slide, cx + pad, gate_y - Inches(0.08), card_w - 2 * pad, 0, color=dk.HAIRLINE)
        gbox = dk.add_textbox(slide, cx + pad, gate_y, card_w - 2 * pad, Inches(0.2))
        gp = gbox.text_frame.paragraphs[0]
        dk.add_run(gp, "THE GATE", font=dk.SANS, size=Pt(9), color=dk.OXBLOOD, bold=True,
                   tracking=1.2)
        dk.add_body_text(slide, cx + pad, gate_y + Inches(0.22), card_w - 2 * pad, Inches(0.5),
                          [[(card["gate"], False)]], size=Pt(10.5), line_spacing=1.15)

    dk.add_footer(slide, deck_name, page)
    return slide


def build_decisions_dark(prs, data, deck_name, page):
    slide = dk.add_blank_slide(prs)
    dk.set_background(slide, dk.DARK_BG)

    y = Inches(0.5)
    dk.add_eyebrow(slide, MARGIN, y, CONTENT_WIDTH,
                   eyebrow_segments(data["breadcrumb"], default_color=dk.TAN_TEXT,
                                     accent_color=dk.TERRACOTTA),
                   with_rule=False)
    y += Inches(0.34)
    dk.add_title(slide, MARGIN, y, CONTENT_WIDTH, Inches(0.7),
                 title_lines(data["title"]), size=dk.SZ_H1,
                 ink_color=dk.CREAM_TEXT, accent_color=dk.TERRACOTTA)
    y += Inches(0.95)

    items = data["items"]
    cols = 2
    col_w = dk.div(CONTENT_WIDTH - Inches(0.6), cols)
    row_h = Inches(1.55)
    for i, item in enumerate(items):
        row, col = divmod(i, cols)
        cx = MARGIN + (col_w + Inches(0.6)) * col
        cy = y + row_h * row
        num_box = dk.add_textbox(slide, cx, cy, Inches(0.7), Inches(0.6))
        p = num_box.text_frame.paragraphs[0]
        dk.add_run(p, str(item.get("number", i + 1)), font=dk.SERIF,
                   size=dk.SZ_DECISION_NUMBER, color=dk.TERRACOTTA, bold=True)
        tx = cx + Inches(0.55)
        tw = col_w - Inches(0.55)
        title_box = dk.add_textbox(slide, tx, cy, tw, Inches(0.3))
        tp = title_box.text_frame.paragraphs[0]
        dk.add_run(tp, item["title"], font=dk.SERIF, size=dk.SZ_DECISION_ITEM_TITLE,
                   color=dk.CREAM_TEXT, bold=True)
        dk.add_body_text(slide, tx, cy + Inches(0.32), tw, Inches(0.75),
                          body_paras([item["body"]]) if isinstance(item["body"], list) else [[(item["body"], False)]],
                          size=Pt(10.5), color=dk.TAN_TEXT, line_spacing=1.2)
        if item.get("cta"):
            dk.add_cta(slide, tx, cy + Inches(1.08), item["cta"], color=dk.TERRACOTTA)

    footer_y = dk.SLIDE_H - Inches(0.45)
    dk.add_line(slide, MARGIN, footer_y - Inches(0.12), CONTENT_WIDTH, 0, color=dk.DARK_HAIRLINE)
    fl = dk.add_textbox(slide, MARGIN, footer_y, Inches(7), Inches(0.25))
    flp = fl.text_frame.paragraphs[0]
    dk.add_run(flp, data.get("footer_left", deck_name).upper(), font=dk.SANS,
               size=dk.SZ_MICRO_FOOTER, color=dk.TAN_TEXT, tracking=1.5)
    fr = dk.add_textbox(slide, dk.SLIDE_W - MARGIN - Inches(2), footer_y, Inches(2), Inches(0.25),
                         align=PP_ALIGN.RIGHT)
    frp = fr.text_frame.paragraphs[0]
    frp.alignment = PP_ALIGN.RIGHT
    dk.add_run(frp, data.get("footer_right", ""), font=dk.SANS, size=dk.SZ_MICRO_FOOTER,
               color=dk.TAN_TEXT, tracking=1.5)
    return slide


BUILDERS = {
    "cover": build_cover,
    "progress": build_progress,
    "two_column": build_two_column,
    "risks": build_risks,
    "decisions_dark": build_decisions_dark,
}


def build(spec_path, output_path):
    with open(spec_path, "r", encoding="utf-8") as f:
        spec = json.load(f)

    deck_name = spec.get("deck_name", "")
    prs = dk.new_presentation()

    for i, slide_spec in enumerate(spec["slides"], start=1):
        stype = slide_spec["type"]
        builder = BUILDERS.get(stype)
        if builder is None:
            raise ValueError(f"Unknown slide type: {stype!r}")
        page = slide_spec.get("page", i)
        if stype == "cover":
            slide_spec = dict(slide_spec)
            slide_spec["page"] = page
            builder(prs, slide_spec, deck_name)
        else:
            builder(prs, slide_spec, deck_name, page)

    prs.save(output_path)
    print(f"Wrote {output_path}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Usage: python3 build_deck.py <content.json> <output.pptx>")
        sys.exit(1)
    build(sys.argv[1], sys.argv[2])
